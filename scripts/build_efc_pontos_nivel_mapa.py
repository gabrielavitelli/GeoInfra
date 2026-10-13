#!/usr/bin/env python3
"""Mapa EFC pontos de nível + estações ANA; PPT slide dos 3 novos pontos."""
from __future__ import annotations

import json
import math
from datetime import datetime
from pathlib import Path

import folium
from folium import plugins
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

MAP_DATA = Path("/tmp/km661_map_data.json")
OUT_DIR = Path("/workspace")
OUT_HTML = OUT_DIR / "EFC_pontos_nivel_mapa.html"
OUT_PNG = OUT_DIR / "EFC_pontos_nivel_mapa.png"
OUT_SUMMARY = OUT_DIR / "EFC_pontos_nivel_resumo.json"
OUT_PPTX = OUT_DIR / "Km661_Ponte_Sumauma_nivel_ultimos_30dias.pptx"
# Keep legacy names as copies for overwrite compatibility
LEGACY_HTML = OUT_DIR / "Km661_Ponte_Sumauma_mapa_estacoes.html"
LEGACY_PNG = OUT_DIR / "Km661_Ponte_Sumauma_mapa_estacoes.png"

# EFC monitoring points (user coords)
EFC_POINTS = [
    {
        "id": "km85",
        "label": "Km 85",
        "subtitle": "Medidor / ponto EFC",
        "codigo": "km85",
        "lat": -3.298651,
        "lon": -44.465063,
        "color": "#1d4ed8",
        "icon": "tint",
        "folium_color": "blue",
        "marker": "o",
    },
    {
        "id": "km374",
        "label": "Sensor Rio Pindaré",
        "subtitle": "KM374 / BC07541S",
        "codigo": "BC07541S",
        "lat": -4.253738,
        "lon": -46.531145,
        "color": "#7c3aed",
        "icon": "tint",
        "folium_color": "purple",
        "marker": "s",
    },
    {
        "id": "km763",
        "label": "Ponte 763+800",
        "subtitle": "BC07533S · Ponte EFC",
        "codigo": "BC07533S",
        "lat": -5.571403,
        "lon": -49.241924,
        "color": "#be185d",
        "icon": "road",
        "folium_color": "pink",
        "marker": "D",
    },
    {
        "id": "km661",
        "label": "Km 661 · Ponte Sumaúma",
        "subtitle": "BQ railway_km_661",
        "codigo": "railway_km_661",
        "lat": -5.149214340420318,
        "lon": -48.37704830878583,
        "color": "#c2410c",
        "icon": "star",
        "folium_color": "red",
        "marker": "*",
    },
]

ANA_ORDER = ["23800495", "29040000", "28850000", "29050000"]
ANA_COLORS = {
    "23800495": "#64748b",
    "29040000": "#0c4a6e",
    "28850000": "#0f766e",
    "29050000": "#b45309",
}


def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlmb = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def nearest_ana(lat: float, lon: float, stations: dict) -> tuple[str, dict, float]:
    best = None
    best_d = 1e18
    for code, s in stations.items():
        d = haversine_km(lat, lon, s["lat"], s["lon"])
        if d < best_d:
            best_d = d
            best = (code, s, d)
    assert best is not None
    return best


def enrich_points(stations: dict) -> list[dict]:
    out = []
    for p in EFC_POINTS:
        code, s, d = nearest_ana(p["lat"], p["lon"], stations)
        item = dict(p)
        item["nearest_ana"] = {
            "codigo": code,
            "nome": s["nome"],
            "dist_km": round(d, 1),
            "lat": s["lat"],
            "lon": s["lon"],
        }
        out.append(item)
    return out


def build_html(data: dict, points: list[dict], path: Path) -> None:
    stations = data["stations"]
    efc = data["efc"]

    # Center roughly on corridor midpoint
    all_lats = [p["lat"] for p in points] + [stations[c]["lat"] for c in ANA_ORDER]
    all_lons = [p["lon"] for p in points] + [stations[c]["lon"] for c in ANA_ORDER]
    m = folium.Map(
        location=[sum(all_lats) / len(all_lats), sum(all_lons) / len(all_lons)],
        zoom_start=6,
        tiles="OpenStreetMap",
        control_scale=True,
    )

    fg_efc = folium.FeatureGroup(name="Pontos EFC (nível)", show=True)
    fg_ana = folium.FeatureGroup(name="Estações ANA (nível)", show=True)
    fg_linha = folium.FeatureGroup(name="EFC km 640–690", show=True)

    efc_coords = [[p["lat"], p["lon"]] for p in efc]
    folium.PolyLine(
        efc_coords,
        color="#1e3a5f",
        weight=4,
        opacity=0.85,
        tooltip="EFC (km 640–690 · trecho Sumaúma)",
    ).add_to(fg_linha)

    for p in points:
        na = p["nearest_ana"]
        html = (
            f"<b>{p['label']}</b><br>"
            f"{p['subtitle']}<br>"
            f"Código: <b>{p['codigo']}</b><br>"
            f"Lat/Lon: {p['lat']:.6f}, {p['lon']:.6f}<br>"
            f"<hr style='margin:6px 0'>"
            f"ANA mais próxima: <b>{na['codigo']}</b><br>"
            f"{na['nome']}<br>"
            f"Distância: <b>{na['dist_km']:.1f} km</b>"
        )
        folium.Marker(
            [p["lat"], p["lon"]],
            popup=folium.Popup(html, max_width=340),
            tooltip=f"{p['label']} · {p['codigo']}",
            icon=folium.Icon(color=p["folium_color"], icon=p["icon"], prefix="fa"),
        ).add_to(fg_efc)
        folium.CircleMarker(
            [p["lat"], p["lon"]],
            radius=8,
            color=p["color"],
            fill=True,
            fill_color=p["color"],
            fill_opacity=0.85,
            weight=2,
        ).add_to(fg_efc)

    for code in ANA_ORDER:
        s = stations[code]
        color = ANA_COLORS[code]
        # distance to Km 661 for continuity with prior report
        km661 = next(p for p in points if p["id"] == "km661")
        d661 = haversine_km(km661["lat"], km661["lon"], s["lat"], s["lon"])
        html = (
            f"<b>{s['nome']}</b><br>"
            f"Código ANA: <b>{code}</b><br>"
            f"Rio: {s.get('rio', '—')}<br>"
            f"{s.get('municipio', '')}<br>"
            f"Lat/Lon: {s['lat']:.5f}, {s['lon']:.5f}<br>"
            f"Distância ao Km 661: <b>{d661:.1f} km</b>"
        )
        folium.CircleMarker(
            [s["lat"], s["lon"]],
            radius=9,
            color=color,
            fill=True,
            fill_color=color,
            fill_opacity=0.9,
            weight=2,
            popup=folium.Popup(html, max_width=320),
            tooltip=f"ANA {code} · {s['nome']}",
        ).add_to(fg_ana)
        folium.Marker(
            [s["lat"], s["lon"]],
            icon=folium.DivIcon(
                html=(
                    f'<div style="font-size:11px;font-weight:700;color:{color};'
                    f'white-space:nowrap;text-shadow:0 0 3px #fff,0 0 3px #fff;">'
                    f"{code}</div>"
                )
            ),
        ).add_to(fg_ana)

    fg_linha.add_to(m)
    fg_efc.add_to(m)
    fg_ana.add_to(m)
    folium.LayerControl(collapsed=False).add_to(m)

    legend_rows = "".join(
        f'<div style="margin-top:3px;"><span style="color:{p["color"]};font-size:14px;">●</span> '
        f'{p["label"]} <span style="color:#64748b;">({p["codigo"]})</span></div>'
        for p in points
    )
    ana_rows = "".join(
        f'<div style="color:{ANA_COLORS[c]};">● ANA {c}</div>' for c in ANA_ORDER
    )
    legend_html = f"""
    <div style="position:fixed;bottom:24px;left:24px;z-index:9999;
                background:rgba(255,255,255,0.96);padding:12px 14px;border-radius:8px;
                border:1px solid #cbd5e1;font-family:Arial,sans-serif;font-size:12px;
                box-shadow:0 2px 8px rgba(0,0,0,.12);max-width:360px;">
      <div style="font-weight:700;margin-bottom:6px;color:#0f172a;">
        EFC · Pontos de nível + ANA
      </div>
      <div style="font-weight:600;color:#334155;margin-bottom:2px;">Pontos EFC</div>
      {legend_rows}
      <div style="font-weight:600;color:#334155;margin:8px 0 2px;">Estações ANA (nível)</div>
      {ana_rows}
      <div style="margin-top:6px;">━ EFC km 640–690 (Sumaúma)</div>
      <div style="margin-top:8px;color:#64748b;font-size:11px;">
        Distâncias em linha reta · popups com ANA mais próxima
      </div>
    </div>
    """
    m.get_root().html.add_child(folium.Element(legend_html))
    m.fit_bounds(
        [[min(all_lats), min(all_lons)], [max(all_lats), max(all_lons)]],
        padding=(50, 50),
    )
    plugins.Fullscreen().add_to(m)
    m.save(str(path))
    print(f"HTML: {path}")


def build_png(data: dict, points: list[dict], path: Path) -> None:
    stations = data["stations"]
    efc = data["efc"]

    fig, ax = plt.subplots(figsize=(13, 9), facecolor="#f8fafc")
    ax.set_facecolor("#eef2f6")

    ax.plot(
        [p["lon"] for p in efc],
        [p["lat"] for p in efc],
        color="#1e3a5f",
        lw=2.2,
        label="EFC km 640–690",
        zorder=2,
    )

    for p in points:
        size = 320 if p["marker"] == "*" else 120
        ax.scatter(
            [p["lon"]],
            [p["lat"]],
            s=size,
            c=p["color"],
            marker=p["marker"],
            zorder=6,
            edgecolors="white",
            linewidths=0.7,
            label=f"{p['label']} ({p['codigo']})",
        )
        ax.annotate(
            f"{p['label']}\n{p['codigo']}",
            (p["lon"], p["lat"]),
            textcoords="offset points",
            xytext=(10, 8),
            fontsize=8,
            fontweight="bold",
            color=p["color"],
        )

    for code in ANA_ORDER:
        s = stations[code]
        color = ANA_COLORS[code]
        ax.scatter(
            [s["lon"]],
            [s["lat"]],
            s=80,
            c=color,
            zorder=5,
            edgecolors="white",
            linewidths=0.8,
            label=f"ANA {code}",
        )
        ax.annotate(
            f"{code}",
            (s["lon"], s["lat"]),
            textcoords="offset points",
            xytext=(6, -10),
            fontsize=7,
            color=color,
            fontweight="bold",
        )

    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_title(
        "EFC — pontos de nível (Km 85, Pindaré, Sumaúma, 763+800) e estações ANA",
        fontsize=12,
        pad=12,
    )
    ax.grid(True, alpha=0.35)
    ax.legend(loc="upper right", fontsize=7.5, framealpha=0.95, ncol=1)
    ax.set_aspect("equal", adjustable="datalim")
    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight")
    plt.close(fig)
    print(f"PNG: {path}")


def _set_run(run, text: str, size: int = 14, bold: bool = False, color=(20, 40, 60)):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = "Calibri"


def add_novos_pontos_slide(prs: Presentation, points: list[dict]) -> None:
    """Slide listing the 3 new points + distance to nearest ANA."""
    novos = [p for p in points if p["id"] in ("km85", "km374", "km763")]
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.95)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(12, 74, 110)
    header.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.55))
    p0 = box.text_frame.paragraphs[0]
    r = p0.add_run()
    _set_run(r, "Novos pontos EFC de nível — localização e ANA mais próxima", 22, True, (255, 255, 255))

    rows = len(novos) + 1
    table = slide.shapes.add_table(
        rows, 6, Inches(0.4), Inches(1.2), Inches(12.5), Inches(0.45 * rows)
    ).table
    headers = ["Ponto", "Código", "Lat", "Lon", "ANA mais próxima", "Dist. (km)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(11)
                run.font.bold = True

    for i, p in enumerate(novos, start=1):
        na = p["nearest_ana"]
        vals = [
            p["label"],
            p["codigo"],
            f"{p['lat']:.6f}",
            f"{p['lon']:.6f}",
            f"{na['codigo']} · {na['nome'][:32]}",
            f"{na['dist_km']:.1f}",
        ]
        for j, v in enumerate(vals):
            cell = table.cell(i, j)
            cell.text = v
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)

    note = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12.3), Inches(3.5))
    tf = note.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    lines = [
        "Detalhamento (coordenadas fornecidas · planilhas xlsx ainda não incorporadas nesta VM)",
        "",
    ]
    for pt in novos:
        na = pt["nearest_ana"]
        lines.append(
            f"• {pt['label']} ({pt['codigo']}): {pt['lat']:.6f}, {pt['lon']:.6f} — "
            f"ANA {na['codigo']} ({na['nome']}) a {na['dist_km']:.1f} km"
        )
    lines += [
        "",
        "• Km 661 Ponte Sumaúma permanece o foco do relatório de nível (30 dias) com as 4 ANA priorizadas.",
        "• Para enriquecer popups com séries das planilhas, enviar os 3 xlsx para "
        "gs://efc-clima-data/reports/km661_ponte_sumauma/entrada/",
        f"• Atualizado em {datetime.now().strftime('%Y-%m-%d %H:%M')}",
    ]
    _set_run(r, "\n".join(lines), 13, False, (40, 60, 80))


def update_pptx(points: list[dict], path: Path) -> None:
    if path.exists():
        prs = Presentation(str(path))
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    add_novos_pontos_slide(prs, points)
    # also add map image slide if PNG exists
    if OUT_PNG.exists():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        _set_run(r, "Mapa atualizado — pontos EFC de nível + ANA", 20, True, (12, 50, 80))
        slide.shapes.add_picture(str(OUT_PNG), Inches(0.55), Inches(0.8), width=Inches(12.2))
    prs.save(str(path))
    print(f"PPTX: {path}")


def main() -> None:
    data = json.loads(MAP_DATA.read_text())
    # ensure km661 uses BQ coords from map data
    for p in EFC_POINTS:
        if p["id"] == "km661":
            p["lat"] = data["km661"]["lat"]
            p["lon"] = data["km661"]["lon"]

    points = enrich_points(data["stations"])
    summary = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "xlsx_incorporados": False,
        "nota": "/home/gabriela e gs://efc-clima-data sem os 3 xlsx; mapa com coordenadas fornecidas.",
        "pontos_efc": [
            {
                "id": p["id"],
                "label": p["label"],
                "codigo": p["codigo"],
                "lat": p["lat"],
                "lon": p["lon"],
                "nearest_ana": p["nearest_ana"],
            }
            for p in points
        ],
        "ana": {c: data["stations"][c] for c in ANA_ORDER},
    }
    OUT_SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"RESUMO: {OUT_SUMMARY}")

    build_html(data, points, OUT_HTML)
    build_png(data, points, OUT_PNG)
    LEGACY_HTML.write_bytes(OUT_HTML.read_bytes())
    LEGACY_PNG.write_bytes(OUT_PNG.read_bytes())
    update_pptx(points, OUT_PPTX)

    for dest in (Path("/opt/cursor/artifacts"), Path("/tmp")):
        dest.mkdir(parents=True, exist_ok=True)
        for f in (OUT_HTML, OUT_PNG, OUT_PPTX, OUT_SUMMARY, LEGACY_HTML, LEGACY_PNG):
            if f.exists():
                (dest / f.name).write_bytes(f.read_bytes())
                print(f"copied → {dest / f.name}")


if __name__ == "__main__":
    main()
