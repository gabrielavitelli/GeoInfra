#!/usr/bin/env python3
"""PPT: situação medidor nível do rio — Km 661 Ponte Sumaúma — últimos 30 dias."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

SUMMARY = Path("/tmp/km661_summary.json")
SERIES = Path("/tmp/km661_series.json")
OUT_PPTX = Path("/workspace/Km661_Ponte_Sumauma_nivel_ultimos_30dias.pptx")
OUT_CHART = Path("/workspace/Km661_Ponte_Sumauma_nivel_chart.png")
OUT_MAP = Path("/tmp/km661_contexto_mapa.png")

# Prefer stations with real variation; keep closest even if flat (flagged).
PRIMARY = ["29040000", "29050000", "28850000", "23800495"]
KM661 = (-5.090, -48.450)


def _set_run(run, text: str, size: int = 14, bold: bool = False, color=(20, 40, 60)):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = "Calibri"


def add_title_slide(prs: Presentation, summary: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(2.2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(12, 74, 110)
    shape.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12), Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "EFC · Monitoramento hidrológico", 16, False, (180, 220, 240))

    box2 = slide.shapes.add_textbox(Inches(0.6), Inches(2.5), Inches(12), Inches(3))
    tf = box2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(
        r,
        "Situação do medidor de nível do rio\nKm 661 — Ponte Sumaúma\nÚltimos 30 dias",
        32,
        True,
        (12, 50, 80),
    )

    box3 = slide.shapes.add_textbox(Inches(0.6), Inches(5.8), Inches(12), Inches(1.5))
    tf = box3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    freshest = max((s["last_ts"] for s in summary.values()), default="—")
    _set_run(
        r,
        f"Fonte: ANA telemétrica (ServiceANA DadosHidrometeorologicos)\n"
        f"Referência espacial: Km 661 EFC ≈ São Pedro da Água Branca/MA "
        f"({KM661[0]:.3f}, {KM661[1]:.3f})\n"
        f"Gerado em {datetime.now().strftime('%Y-%m-%d %H:%M')} · última obs. ANA nas estações: {freshest}",
        13,
        False,
        (60, 80, 100),
    )


def make_context_map(summary: dict, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 6.2), facecolor="#f4f7f9")
    ax.set_facecolor("#e8eef2")
    ax.scatter([KM661[1]], [KM661[0]], s=220, c="#c2410c", zorder=5, marker="*", label="Km 661 (aprox.)")
    colors = ["#0c4a6e", "#0369a1", "#0e7490", "#64748b"]
    for i, code in enumerate(PRIMARY):
        if code not in summary:
            continue
        s = summary[code]
        ax.scatter([s["lon"]], [s["lat"]], s=90, c=colors[i % len(colors)], zorder=4, label=f"{code} · {s['dist_km']} km")
        ax.annotate(
            f"{s['nome']}\n{code}",
            (s["lon"], s["lat"]),
            textcoords="offset points",
            xytext=(8, 6),
            fontsize=8,
            color="#0f172a",
        )
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_title("Contexto espacial — Km 661 e estações ANA (nível) mais próximas")
    ax.grid(True, alpha=0.35)
    ax.legend(loc="best", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def make_chart(series: dict, summary: dict, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11.5, 5.8), facecolor="white")
    colors = ["#0c4a6e", "#b45309", "#0f766e", "#64748b"]
    for i, code in enumerate(PRIMARY):
        if code not in series:
            continue
        pts = series[code]["nivel"]
        if not pts:
            continue
        # daily mean in meters (ANA Nivel em cm)
        by_day = {}
        for ts, cm in pts:
            d = ts[:10]
            by_day.setdefault(d, []).append(cm / 100.0)
        days = sorted(by_day)
        vals = [sum(by_day[d]) / len(by_day[d]) for d in days]
        xs = [datetime.strptime(d, "%Y-%m-%d") for d in days]
        label = f"{summary[code]['nome']} ({code})"
        if code == "23800495":
            label += " [cota constante — checar sensor]"
        ax.plot(xs, vals, color=colors[i % len(colors)], lw=2.0, marker="o", ms=3, label=label)
    ax.set_ylabel("Nível / cota (m) — ANA Nivel÷100")
    ax.set_xlabel("Data")
    ax.set_title("Nível do rio — estações ANA mais próximas do Km 661 (últimos ~30 dias)")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m"))
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8, loc="best")
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(path, dpi=170)
    plt.close(fig)


def add_image_slide(prs: Presentation, title: str, img: Path, note: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(r, title, 22, True, (12, 50, 80))
    slide.shapes.add_picture(str(img), Inches(0.45), Inches(0.95), width=Inches(12.4))
    if note:
        box2 = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12), Inches(0.5))
        p = box2.text_frame.paragraphs[0]
        r = p.add_run()
        _set_run(r, note, 11, False, (70, 90, 110))


def add_table_slide(prs: Presentation, summary: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(r, "Estatísticas de nível (últimos ~30 dias)", 22, True, (12, 50, 80))

    rows = [c for c in PRIMARY if c in summary]
    table = slide.shapes.add_table(len(rows) + 1, 9, Inches(0.3), Inches(1.0), Inches(12.7), Inches(0.4 * (len(rows) + 1))).table
    headers = [
        "Código",
        "Estação",
        "Dist. (km)",
        "Min (m)",
        "Máx (m)",
        "Média (m)",
        "Último (m)",
        "Última obs.",
        "Vazão últ.",
    ]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.bold = True

    for i, code in enumerate(rows, start=1):
        s = summary[code]
        vals = [
            code,
            s["nome"][:28],
            f"{s['dist_km']:.1f}",
            f"{s['min_cm']/100:.2f}",
            f"{s['max_cm']/100:.2f}",
            f"{s['mean_cm']/100:.2f}",
            f"{s['last_cm']/100:.2f}",
            s["last_ts"][5:16],
            f"{s['last_vazao']:.1f}" if s.get("last_vazao") is not None else "—",
        ]
        for j, v in enumerate(vals):
            cell = table.cell(i, j)
            cell.text = v
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    note = slide.shapes.add_textbox(Inches(0.5), Inches(5.2), Inches(12), Inches(2))
    tf = note.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(
        r,
        "Notas operacionais\n"
        "• Não há estação ANA exatamente no Km 661 / Ponte Sumaúma; usamos as telemétricas com nível "
        "mais próximas da referência EFC (~São Pedro da Água Branca).\n"
        "• Cota ANA (campo Nivel) convertida para metros (÷100). Vazão em m³/s quando disponível.\n"
        "• 23800495 (São Sebastião do TO) está ~33 km, mas a cota veio constante (230 cm) no período — "
        "tratar com cautela.\n"
        "• Referências úteis para ops: 29040000 (UHE Tucuruí Montante 1) e 29050000 (Marabá).\n"
        "• BigQuery climaticsystem.efc_clima.clima_series (source=ana, variable=nivel_m) estava defasado "
        "em 2026-09-15; a API pública ANA já responde até hoje — ingestão forçada recomendada.",
        12,
        False,
        (40, 60, 80),
    )


def add_stations_slide(prs: Presentation, summary: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(r, "Estações ANA mais próximas — identificação", 22, True, (12, 50, 80))

    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for code in PRIMARY:
        if code not in summary:
            continue
        s = summary[code]
        if not first:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]
            first = False
        r = p.add_run()
        _set_run(
            r,
            f"{code} — {s['nome']}\n"
            f"   Rio: {s['rio']} · {s['municipio']}\n"
            f"   Lat/Lon: {s['lat']:.5f}, {s['lon']:.5f} · Distância ao Km 661 ≈ {s['dist_km']:.1f} km\n"
            f"   Última obs.: {s['last_ts']} · nível {s['last_cm']/100:.2f} m"
            + (f" · vazão {s['last_vazao']:.1f} m³/s" if s.get("last_vazao") is not None else ""),
            14,
            False,
            (25, 45, 65),
        )
        p.space_after = Pt(10)


def main() -> None:
    summary = json.loads(SUMMARY.read_text())
    series = json.loads(SERIES.read_text())
    make_context_map(summary, OUT_MAP)
    make_chart(series, summary, OUT_CHART)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_title_slide(prs, summary)
    add_stations_slide(prs, summary)
    add_image_slide(
        prs,
        "Mapa / contexto — Km 661 e estações ANA",
        OUT_MAP,
        "Referência Km 661 aproximada a partir do rotograma EFC (São Pedro ~km 658,7) + município OSM.",
    )
    add_image_slide(
        prs,
        "Série temporal de nível (m) — média diária",
        OUT_CHART,
        "Fonte ANA telemétrica pública · valores em metros (Nivel cm ÷ 100).",
    )
    add_table_slide(prs, summary)
    prs.save(OUT_PPTX)
    print(f"PPTX: {OUT_PPTX}")
    print(f"PNG:  {OUT_CHART}")


if __name__ == "__main__":
    main()
